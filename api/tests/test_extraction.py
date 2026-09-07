"""Tests for app.ingestion.extraction. PDF, DOCX, TXT extraction.

DOCX and TXT cases use real files built in ``tmp_path`` (python-docx for the
documents, hand-encoded bytes for the text files) because the defects they pin
were all in what the loaders did NOT read: tables, headers and footers dropped
from DOCX, and non-UTF-8 text silently destroyed by a hard-coded decode.
"""

import codecs
import io
from unittest.mock import MagicMock, mock_open, patch

import docx
import pytest
from pypdf import PasswordType, PdfReader, PdfWriter
from pypdf.errors import DependencyError

from app.ingestion.extraction import ExtractionError, decode_text_bytes, load_docx, load_pdf, load_txt

# ── PDF ─────────────────────────────────────────────────────────────────────


def _mock_reader(pages, *, encrypted: bool = False):
    reader = MagicMock()
    reader.pages = pages
    reader.is_encrypted = encrypted
    return reader


def _page(text):
    page = MagicMock()
    page.extract_text.return_value = text
    return page


def _pdf_with_text(text: str) -> bytes:
    """A minimal one-page PDF whose content stream draws ``text`` in Helvetica.

    pypdf can write PDFs but cannot draw text, so the objects are laid out by
    hand with a correct xref table.
    """
    content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode())
    return out.getvalue()


def _encrypted_pdf(path, *, user_password: str) -> None:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_pdf_with_text("Quarterly report"))))
    writer.encrypt(user_password=user_password, owner_password="owner-secret", algorithm="AES-128")
    with open(path, "wb") as f:
        writer.write(f)


class TestLoadPdf:
    def test_extracts_pages_with_metadata(self):
        reader = _mock_reader([_page("Page one text"), _page("Page two text")])

        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", return_value=reader),
        ):
            result = load_pdf("test.pdf")

        assert len(result) == 2
        assert result[0]["text"] == "Page one text"
        assert result[0]["metadata"]["page"] == 1
        assert result[0]["metadata"]["total_pages"] == 2
        assert result[1]["metadata"]["page"] == 2

    def test_skips_pages_with_no_text(self):
        reader = _mock_reader([_page("Has text"), _page(None), _page("")])

        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", return_value=reader),
        ):
            result = load_pdf("test.pdf")

        assert len(result) == 1
        assert result[0]["text"] == "Has text"
        assert result[0]["metadata"]["total_pages"] == 3

    def test_empty_pdf_raises(self):
        """A PDF with no extractable text (e.g. scanned) must raise so the upload
        route can surface a clear message instead of silently storing zero chunks.
        """
        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", return_value=_mock_reader([])),
            pytest.raises(ExtractionError, match="No extractable text"),
        ):
            load_pdf("empty.pdf")

    def test_scanned_pdf_raises(self):
        """All pages returning empty text (typical scanned PDF) must raise."""
        reader = _mock_reader([_page(""), _page(None)])

        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", return_value=reader),
            pytest.raises(ExtractionError, match="scanned"),
        ):
            load_pdf("scanned.pdf")

    def test_real_unencrypted_pdf_extracts_its_text(self, tmp_path):
        path = tmp_path / "plain.pdf"
        path.write_bytes(_pdf_with_text("Quarterly report"))

        result = load_pdf(str(path))

        assert "Quarterly report" in result[0]["text"]

    def test_password_protected_pdf_raises_a_clear_message(self, tmp_path):
        """pypdf raises ``FileNotDecryptedError`` on the first ``.pages`` access,
        which the upload UI rendered as a generic failure. The customer must be
        told what is wrong with the file and what to do about it."""
        path = tmp_path / "locked.pdf"
        _encrypted_pdf(path, user_password="secret")

        with pytest.raises(ExtractionError, match="password-protected"):
            load_pdf(str(path))

    def test_owner_locked_pdf_opens_with_the_empty_password(self, tmp_path):
        """Print/copy-restricted PDFs are "encrypted" with an empty user
        password. They are the common case and must extract normally."""
        path = tmp_path / "owner-locked.pdf"
        _encrypted_pdf(path, user_password="")

        result = load_pdf(str(path))

        assert "Quarterly report" in result[0]["text"]

    def test_decrypt_is_attempted_before_reading_pages(self):
        reader = _mock_reader([_page("Unlocked text")], encrypted=True)
        reader.decrypt.return_value = PasswordType.OWNER_PASSWORD

        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", return_value=reader),
        ):
            result = load_pdf("owner.pdf")

        reader.decrypt.assert_called_once_with("")
        assert result[0]["text"] == "Unlocked text"

    def test_missing_crypto_dependency_is_reported_as_a_server_problem(self):
        """Blaming the customer's file for a broken server install would send
        them re-saving PDFs for nothing."""
        reader = _mock_reader([], encrypted=True)
        reader.decrypt.side_effect = DependencyError("cryptography>=3.1 is required for AES algorithm")

        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", return_value=reader),
            pytest.raises(ExtractionError, match="server is missing"),
        ):
            load_pdf("aes.pdf")


# ── DOCX ────────────────────────────────────────────────────────────────────


def _save(doc, tmp_path, name="doc.docx") -> str:
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def _lines(result) -> list[str]:
    assert len(result) == 1
    assert result[0]["metadata"] == {"page": 1, "total_pages": 1}
    return result[0]["text"].split("\n")


class TestLoadDocx:
    def test_paragraph_only_document_joins_paragraphs_with_newlines(self, tmp_path):
        """The pre-existing contract: paragraphs joined by ``\\n``, empty ones
        skipped. It must hold so re-uploading an unchanged file still dedups."""
        doc = docx.Document()
        doc.add_paragraph("Para 1")
        doc.add_paragraph("")
        doc.add_paragraph("Para 2")

        assert _lines(load_docx(_save(doc, tmp_path))) == ["Para 1", "Para 2"]

    def test_tables_are_emitted_as_pipe_rows_in_document_order(self, tmp_path):
        """``doc.paragraphs`` dropped every table, so a pricing sheet laid out
        as a table uploaded as an empty-looking document. The rows must land
        between the paragraphs that introduce and discuss the table."""
        doc = docx.Document()
        doc.add_paragraph("Our plans")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Plan"
        table.cell(0, 1).text = "Price"
        table.cell(1, 0).text = "Pro"
        table.cell(1, 1).text = "$99 per\nmonth"  # a multi-paragraph cell stays on its row
        doc.add_paragraph("Contact sales for volume pricing.")

        assert _lines(load_docx(_save(doc, tmp_path))) == [
            "Our plans",
            "| Plan | Price |",
            "| Pro | $99 per month |",
            "Contact sales for volume pricing.",
        ]

    def test_a_merged_cell_is_emitted_once(self, tmp_path):
        """Word reports a horizontally merged cell once per grid column it
        spans; the row must not repeat it."""
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=3)
        table.cell(0, 0).text = "Enterprise"
        table.cell(0, 2).text = "Custom"
        table.cell(0, 0).merge(table.cell(0, 1))

        assert _lines(load_docx(_save(doc, tmp_path))) == ["| Enterprise | Custom |"]

    def test_headers_and_footers_are_appended_once(self, tmp_path):
        doc = docx.Document()
        doc.add_paragraph("Body")
        first = doc.sections[0]
        first.header.paragraphs[0].text = "Acme Confidential"
        first.footer.paragraphs[0].text = "Version 2.1"
        # A second section linked to the first inherits the same header and
        # footer; they must not be emitted twice.
        second = doc.add_section()
        assert second.header.is_linked_to_previous
        doc.add_paragraph("More body")

        assert _lines(load_docx(_save(doc, tmp_path))) == ["Body", "More body", "Acme Confidential", "Version 2.1"]

    def test_a_document_without_headers_reads_cleanly(self, tmp_path):
        doc = docx.Document()
        doc.add_paragraph("Only body")

        assert _lines(load_docx(_save(doc, tmp_path))) == ["Only body"]

    def test_empty_document_raises(self, tmp_path):
        """A DOCX with no text anywhere must raise rather than silently store a
        blank 'document'."""
        doc = docx.Document()
        doc.add_paragraph("")
        doc.add_table(rows=1, cols=1)

        with pytest.raises(ExtractionError, match="empty or image-only"):
            load_docx(_save(doc, tmp_path))


# ── TXT / MD ────────────────────────────────────────────────────────────────


def _txt(tmp_path, raw: bytes, name="notes.txt") -> str:
    path = tmp_path / name
    path.write_bytes(raw)
    return str(path)


class TestLoadTxt:
    def test_reads_utf8_verbatim(self, tmp_path):
        content = "Hello world\nLine two"
        result = load_txt(_txt(tmp_path, content.encode("utf-8")))

        assert len(result) == 1
        assert result[0]["text"] == content
        assert result[0]["metadata"] == {"page": 1, "total_pages": 1}

    def test_utf8_bom_is_consumed(self, tmp_path):
        result = load_txt(_txt(tmp_path, codecs.BOM_UTF8 + "Café".encode()))

        assert result[0]["text"] == "Café"

    def test_cp1252_accents_and_smart_quotes_survive(self, tmp_path):
        """Excel and older Word save text as Windows-1252. The hard-coded UTF-8
        read with ``errors="ignore"`` deleted every accented letter and curly
        quote from such files."""
        raw = "Café — “quoted” résumé".encode("cp1252")
        result = load_txt(_txt(tmp_path, raw))

        assert result[0]["text"] == "Café — “quoted” résumé"

    def test_utf16_with_bom_decodes(self, tmp_path):
        """Notepad's "Unicode" encoding. Read as UTF-8 it became a run of NULs."""
        text = "Café — résumé\nsecond line"
        little = load_txt(_txt(tmp_path, text.encode("utf-16"), name="le.txt"))
        big = load_txt(_txt(tmp_path, codecs.BOM_UTF16_BE + text.encode("utf-16-be"), name="be.txt"))

        assert little[0]["text"] == text
        assert big[0]["text"] == text

    def test_line_endings_are_normalised_like_the_old_text_mode_read(self, tmp_path):
        result = load_txt(_txt(tmp_path, b"alpha\r\nbeta\rgamma"))

        assert result[0]["text"] == "alpha\nbeta\ngamma"

    def test_undecodable_bytes_are_marked_not_dropped(self, tmp_path):
        """0x81 is invalid UTF-8 and undefined in cp1252: the last-resort decode
        must leave a visible marker rather than silently join the words."""
        result = load_txt(_txt(tmp_path, b"ok \x81 end"))

        assert result[0]["text"] == "ok � end"

    def test_empty_file_raises(self, tmp_path):
        """An empty TXT must raise so the caller can surface a clear error
        instead of silently storing a blank 'document'."""
        with pytest.raises(ExtractionError, match="empty"):
            load_txt(_txt(tmp_path, b"   \n\t"))


class TestDecodeTextBytes:
    def test_utf32_bom_is_honoured_before_utf16(self):
        # The UTF-32-LE BOM begins with the UTF-16-LE BOM bytes; decoding it
        # as UTF-16 would interleave NULs into every character.
        assert decode_text_bytes("Café".encode("utf-32")) == "Café"

    def test_plain_ascii_is_unchanged(self):
        assert decode_text_bytes(b"plain ascii") == "plain ascii"


class TestPdfConstructorFailures:
    """pypdf reads /Encrypt and tries the empty password inside
    ``PdfReader.__init__``, so a missing crypto dependency or an unsupported
    scheme raises from the constructor, before ``_unlock_pdf`` ever runs. Those
    must reach the customer as actionable messages too."""

    def test_missing_crypto_dependency_on_construction(self):
        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", side_effect=DependencyError("cryptography>=3.1 is required")),
            pytest.raises(ExtractionError, match="missing the library"),
        ):
            load_pdf("x.pdf")

    def test_unsupported_scheme_on_construction(self):
        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", side_effect=NotImplementedError("Unsupported /V value")),
            pytest.raises(ExtractionError, match="encryption scheme"),
        ):
            load_pdf("x.pdf")

    def test_a_corrupt_file_keeps_its_original_error(self):
        # A constructor PdfReadError is a corrupt or non-PDF file, never an
        # encryption problem: the generic handling upstream must still see
        # pypdf's own exception, not a misleading encryption message.
        from pypdf.errors import PdfReadError

        with (
            patch("builtins.open", mock_open()),
            patch("app.ingestion.extraction.PdfReader", side_effect=PdfReadError("EOF marker not found")),
            pytest.raises(PdfReadError),
        ):
            load_pdf("x.pdf")
